#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
每日油价推送 + 预警
1. 汇总表：去掉国际原油，国内/香港/越南显示下次调整估算
2. 预警：越南柴油破 25000 + 香港柴油破 37 HKD
"""

import os
import csv
import json
import smtplib
import requests
import openpyxl
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.header import Header
from datetime import datetime, timedelta, timezone

# ==================== 配置 ====================
MAIL_USERNAME = os.environ.get("MAIL_USERNAME")
MAIL_PASSWORD = os.environ.get("MAIL_PASSWORD")
SMTP_HOST = "smtp.163.com"
SMTP_PORT = 465

def parse_emails(env_name):
    raw = os.environ.get(env_name, "")
    return [e.strip() for e in raw.split(",") if e.strip()]

DAILY_RECEIVERS = parse_emails("DAILY_RECEIVERS")
ALERT_RECEIVERS = parse_emails("ALERT_RECEIVERS")

BASE_URL = "https://no532.github.io/Oil-price-update"
LOCAL_BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

VN_DIESEL_THRESHOLD = 25000
HK_DIESEL_THRESHOLD = 37.0
ALERT_STATE_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "alert_state.json"
)

# ==================== 时间 ====================
def bj_now():
    return datetime.now(timezone(timedelta(hours=8)))

# ==================== 工具 ====================
def http_get_json(url, timeout=15):
    try:
        r = requests.get(url, timeout=timeout, headers={"User-Agent": "Mozilla/5.0"})
        r.raise_for_status()
        return r.json()
    except: return None

def load_json_local(path):
    try:
        with open(path, encoding="utf-8") as f: return json.load(f)
    except: return None

def load_json(url, local):
    return http_get_json(url) or load_json_local(local)

def round2(v):
    return None if v is None else round(v * 100) / 100

def fmt_num(v):
    return "—" if v is None else f"{float(v):.2f}"

def fmt_int(v):
    return "—" if v is None else f"{int(v):,}"

def color_of(v):
    if v is None: return "#8492a6"
    try:
        f = float(v)
        return "#e74c3c" if f > 0 else ("#27ae60" if f < 0 else "#8492a6")
    except: return "#8492a6"

def parse_excel_date(v):
    if isinstance(v, datetime): return v
    if isinstance(v, str):
        for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d"):
            try: return datetime.strptime(v.strip(), fmt)
            except: pass
    return None

# ==================== 数据加载 ====================
def load_price_calendar():
    path = os.path.join(LOCAL_BASE, "国内", "油价调整日历.xlsx")
    if not os.path.exists(path): return []
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Sheet1"]
        dates = []
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=1).value
            if v:
                d = parse_excel_date(v)
                if d: dates.append(d)
        dates.sort()
        return dates
    except: return []

def load_province_history():
    path = os.path.join(LOCAL_BASE, "国内", "各省历史油价数据.xlsx")
    if not os.path.exists(path): return {}
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        result = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in range(2, ws.max_row + 1):
                d = parse_excel_date(ws.cell(row=row, column=1).value)
                v = ws.cell(row=row, column=6).value
                if d and v is not None:
                    try: rows.append((d, float(v)))
                    except: pass
            if rows:
                rows.sort(key=lambda x: x[0], reverse=True)
                result[sheet_name] = rows
        return result
    except: return {}

def load_hk_history():
    path = os.path.join(LOCAL_BASE, "香港", "香港历史油价.csv")
    if not os.path.exists(path): return []
    try:
        with open(path, encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            rows = []
            for r in reader:
                d = r.get("Date", "").strip()
                try: dt = datetime.strptime(d, "%Y/%m/%d")
                except: continue
                vals = [float(r[c]) for c in ["中石化 (零售牌價)", "中國石油 (零售牌價)", "加德士 (零售牌價)", "埃索 (零售牌價)", "蜆殼 (零售牌價)"] if r.get(c) and r[c].replace('.','',1).isdigit()]
                if vals: rows.append((dt, sum(vals) / len(vals)))
            rows.sort(key=lambda x: x[0], reverse=True)
            return rows
    except: return []

def load_vn_history():
    path = os.path.join(LOCAL_BASE, "越南", "越南历史油价.csv")
    if not os.path.exists(path): return []
    try:
        with open(path, encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            rows = []
            for r in reader:
                d = r.get("Date", "").strip()
                v = r.get("DO005S_M2_Z1", "").strip()
                if d and v:
                    try: rows.append((datetime.strptime(d, "%Y/%m/%d"), float(v)))
                    except: pass
            rows.sort(key=lambda x: x[0], reverse=True)
            return rows
    except: return []

def load_intl_monthly():
    d = load_json(f"{BASE_URL}/国际/oil_monthly.json", os.path.join(LOCAL_BASE, "国际", "oil_monthly.json"))
    if not d: return {}
    result = {}
    for name, arr in (d.get("series") or {}).items():
        rows = []
        for item in arr:
            try: rows.append((datetime.strptime(item[0], "%Y-%m-%d"), float(item[1])))
            except: pass
        rows.sort(key=lambda x: x[0], reverse=True)
        result[name] = rows
    return result

def fetch_domestic():
    d = load_json(f"{BASE_URL}/国内/data.json", os.path.join(LOCAL_BASE, "国内", "data.json"))
    if not d: return None
    
    rows = []
    for k, v in d.items():
        # 安全检查：必须是列表，且长度大于 0，且第一项必须是字典
        if not isinstance(v, list) or len(v) == 0:
            continue
        latest = v[0]
        if not isinstance(latest, dict):
            continue
            
        rows.append({
            "province": k, 
            "date": latest.get("date", "—"), 
            "p92": latest.get("92号汽油"), 
            "p95": latest.get("95号汽油"), 
            "p0": latest.get("0号柴油"), 
            "d0": latest.get("0号柴油_diff")
        })
        
    rows.sort(key=lambda x: x["province"])
    p0_vals = [r["p0"] for r in rows if r["p0"] is not None]
    avg_p0 = sum(p0_vals)/len(p0_vals) if p0_vals else None
    return {"rows": rows, "avg_p0": avg_p0}

def fetch_hongkong():
    d = load_json(f"{BASE_URL}/香港/hk_fuel.json", os.path.join(LOCAL_BASE, "香港", "hk_fuel.json"))
    if not d: return None
    open_data = d.get("open_data", {})
    rows = []
    for fuel_name, brands in open_data.items():
        if isinstance(brands, dict):
            vals = [v for v in brands.values() if isinstance(v, (int, float))]
            rows.append({"fuel": fuel_name, "retail": sum(vals)/len(vals) if vals else None})
    return {"date": d.get("date", "—"), "update_time": d.get("update_time", "—"), "rows": rows}

def fetch_international():
    d = load_json(f"{BASE_URL}/国际/oil_minline.json", os.path.join(LOCAL_BASE, "国际", "oil_minline.json"))
    if not d: return None
    rows = []
    for name, pts in (d.get("series") or {}).items():
        if pts:
            open_p = (d.get("meta") or {}).get(name, {}).get("open")
            rows.append({"name": name, "cur": open_p if open_p else pts[-1][1]})
    return {"update_time": d.get("update_time", "—"), "rows": rows}

def fetch_vietnam():
    d = load_json(f"{BASE_URL}/越南/vietnam_fuel.json", os.path.join(LOCAL_BASE, "越南", "vietnam_fuel.json"))
    if not d: return None
    order = ["DO005S_M2","DO001S_M5","E10RON95_M5","E10RON95_M3","E5RON92_M2","Kero2K"]
    return {"date": d.get("date", "—"), "update_time": d.get("update_time", "—"), "rows": [{"prefix": k, "name": d.get("fuel_names", {}).get(k, k), "z1": d.get("prices", {}).get(k, {}).get("Z1"), "z2": d.get("prices", {}).get(k, {}).get("Z2")} for k in order if d.get("prices", {}).get(k)]}

# ==================== 预警与估算 ====================
def estimate_hk_next_date(hk_history):
    """根据历史记录估算香港下次调价日期"""
    if not hk_history or len(hk_history) < 2: return "—"
    # 取最近 30 次价格变动的间隔天数
    change_dates = []
    last_p = hk_history[0][1]
    for d, p in hk_history:
        if abs(p - last_p) > 0.01:
            change_dates.append(d)
            last_p = p
            if len(change_dates) >= 30: break
    if len(change_dates) < 2: return "—"
    
    # 计算平均间隔
    intervals = [(change_dates[i] - change_dates[i+1]).days for i in range(len(change_dates)-1)]
    avg_days = sum(intervals) / len(intervals) if intervals else 7
    last_change = change_dates[0]
    
    # 估算下一次
    next_date = last_change + timedelta(days=int(round(avg_days)))
    return f"约 {next_date.strftime('%Y-%m-%d')}（推算）"

def check_alerts(cur_vn, cur_hk, state, today):
    alerts = []
    # 1. 越南预警
    if cur_vn is not None:
        vn_below = cur_vn < VN_DIESEL_THRESHOLD
        vn_was_below = state.get("vn_below", False)
        if vn_below and not vn_was_below: alerts.append("trigger_vn")
        if not vn_below and vn_was_below: alerts.append("recover_vn")
    
    # 2. 香港预警 (仅在下跌破37、或上涨破37时)
    if cur_hk is not None:
        hk_below = cur_hk < HK_DIESEL_THRESHOLD
        hk_was_below = state.get("hk_below", False)
        if hk_below and not hk_was_below: alerts.append("trigger_hk")
        if not hk_below and hk_was_below: alerts.append("recover_hk")
        
    return alerts

# ==================== 组装邮件 ====================
def build_summary(domestic, hk, vn, prov_history, hk_history, vn_history, calendar):
    today = bj_now()
    rows = []

    # 国内
    if domestic and domestic["avg_p0"] is not None:
        avg = domestic["avg_p0"]
        d0_vals = [r["d0"] for r in domestic["rows"] if r["d0"] is not None]
        avg_d0 = sum(d0_vals) / len(d0_vals) if d0_vals else 0
        diff = round2(avg_d0); pct = round2(diff / (avg - diff) * 100) if avg and diff else 0.0
        past = [d for d in calendar if d.date() <= today.date()]
        last_date_str = past[-2].strftime('%Y-%m-%d') if len(past) >= 2 else (past[-1].strftime('%Y-%m-%d') if past else "—")
        next_d = next((d for d in calendar if d.date() >= today.date()), None)
        next_str = f"{next_d.strftime('%Y-%m-%d')}（还有 {(next_d.date()-today.date()).days} 天）" if next_d else "—"
        rows.append({"label": "国内柴油（31 省平均）", "value": f"{fmt_num(avg)}<span class='u'>元/升</span>", "last_date": last_date_str, "diff": diff, "pct": pct, "next": next_str})

    # 香港
    if hk:
        for r in hk["rows"]:
            if "柴油" in r["fuel"]:
                cur = r["retail"]; last_date, last_price = None, None
                for d, p in hk_history:
                    if abs(p - cur) > 0.01: last_date, last_price = d, p; break
                diff = cur - last_price if last_date else None; pct = diff / last_price * 100 if diff is not None else None
                last_date_str = last_date.strftime('%Y-%m-%d') if last_date else "—"
                next_str = estimate_hk_next_date(hk_history)
                rows.append({"label": "香港柴油（平均零售价）", "value": f"{fmt_num(cur)}<span class='u'>HKD/升</span>", "last_date": last_date_str, "diff": diff, "pct": pct, "next": next_str})
                break

    # 越南
    if vn:
        for r in vn["rows"]:
            if r["prefix"] == "DO005S_M2":
                cur = r["z1"]; last_date, last_price = None, None
                for d, p in vn_history:
                    if abs(p - cur) > 0.01: last_date, last_price = d, p; break
                diff = cur - last_price if last_date else None; pct = diff / last_price * 100 if diff is not None else None
                last_date_str = last_date.strftime('%Y-%m-%d') if last_date else "—"
                next_str = "每 7 天调价" # 越南通常7天调一次
                rows.append({"label": "越南柴油 0.05S（区域 1）", "value": f"{fmt_int(cur)}<span class='u'>VND/升</span>", "last_date": last_date_str, "diff": diff, "pct": pct, "next": next_str})
                break

    html = """<div style="background:#fff8f7;border:1px solid #f5d0cc;border-radius:8px;padding:16px 20px;margin-bottom:20px;"><h3 style="color:#c9302c;margin:0 0 12px;font-size:16px;">📊 今日汇总</h3><table style="width:100%;border-collapse:collapse;font-size:13px;"><thead><tr style="background:#fbeae8;"><th style="padding:8px 10px;text-align:left;color:#c9302c;border-bottom:1px solid #f0d0cc;">品种</th><th style="padding:8px 10px;text-align:right;color:#c9302c;border-bottom:1px solid #f0d0cc;">当前价</th><th style="padding:8px 10px;text-align:center;color:#c9302c;border-bottom:1px solid #f0d0cc;">较上次</th><th style="padding:8px 10px;text-align:center;color:#c9302c;border-bottom:1px solid #f0d0cc;">下次调整</th></tr></thead><tbody>"""
    for r in rows:
        if r["diff"] is not None and r["pct"] is not None:
            c = color_of(r["diff"]); sign = "+" if r["diff"] > 0 else ""
            diff_str = f"<div style='color:{c};font-weight:600;'>{sign}{fmt_num(r['diff'])}（{sign}{r['pct']:.2f}%）</div><div style='color:#8492a6;font-size:11px;margin-top:2px;'>较 {r['last_date']}</div>"
        else:
            diff_str = f"<div style='color:#8492a6;'>—</div><div style='color:#8492a6;font-size:11px;margin-top:2px;'>较 —</div>"
        html += f"""<tr><td style="padding:10px;text-align:left;border-bottom:1px dashed #f0d0cc;color:#5a6b7b;">{r['label']}</td><td style="padding:10px;text-align:right;border-bottom:1px dashed #f0d0cc;font-weight:600;color:#1f2d3d;">{r['value']}</td><td style="padding:10px;text-align:center;border-bottom:1px dashed #f0d0cc;">{diff_str}</td><td style="padding:10px;text-align:center;border-bottom:1px dashed #f0d0cc;color:#5a6b7b;font-size:12px;">{r['next']}</td></tr>"""
    return html + "</tbody></table><style>.u { font-size:11px; color:#8492a6; font-weight:400; margin-left:2px; }</style></div>"

def section(title, content):
    return f"""<div style="margin:24px 0;"><h3 style="color:#c9302c;border-left:4px solid #c9302c;padding-left:10px;margin:0 0 10px;font-size:15px;">{title}</h3>{content}</div>"""

def render_domestic(d):
    if not d: return ""
    html = f"""<p style="color:#8492a6;font-size:12px;">共 {len(d['rows'])} 个省市</p><table border="1" cellpadding="8" cellspacing="0" style="border-collapse:collapse;font-size:13px;width:100%;"><thead style="background:#fafbfc;"><tr><th align="left">省份</th><th>日期</th><th>92# (元/升)</th><th>95# (元/升)</th><th>0# 柴油 (元/升)</th></tr></thead><tbody>"""
    for r in d["rows"]: html += f"<tr><td>{r['province']}</td><td align='center'>{r['date']}</td><td align='right'>{fmt_num(r['p92'])}</td><td align='right'>{fmt_num(r['p95'])}</td><td align='right'>{fmt_num(r['p0'])}</td></tr>"
    return html + "</tbody></table>"

def render_hongkong(d):
    if not d: return ""
    html = f"""<p style="color:#8492a6;font-size:12px;">数据日期：{d['date']}　更新时间：{d['update_time']}</p><table border="1" cellpadding="8" cellspacing="0" style="border-collapse:collapse;font-size:13px;width:100%;"><thead style="background:#fafbfc;"><tr><th align="left">油品</th><th>平均零售价 (HKD/L)</th></tr></thead><tbody>"""
    for r in d["rows"]: html += f"<tr><td>{r['fuel']}</td><td align='right'>{fmt_num(r['retail'])}</td></tr>"
    return html + "</tbody></table>"

def render_international(d):
    if not d: return ""
    html = f"""<p style="color:#8492a6;font-size:12px;">更新时间：{d['update_time']}</p><table border="1" cellpadding="8" cellspacing="0" style="border-collapse:collapse;font-size:13px;width:100%;"><thead style="background:#fafbfc;"><tr><th align="left">品种</th><th>今开 (美元/桶)</th></tr></thead><tbody>"""
    for r in d["rows"]: html += f"<tr><td>{r['name']}</td><td align='right'>{fmt_num(r['cur'])}</td></tr>"
    return html + "</tbody></table>"

def render_vietnam(d):
    if not d: return ""
    html = f"""<p style="color:#8492a6;font-size:12px;">数据日期：{d['date']}　更新时间：{d['update_time']}</p><table border="1" cellpadding="8" cellspacing="0" style="border-collapse:collapse;font-size:13px;width:100%;"><thead style="background:#fafbfc;"><tr><th align="left">油品</th><th>区域 1 (VND/L)</th><th>区域 2 (VND/L)</th></tr></thead><tbody>"""
    for r in d["rows"]: html += f"<tr><td>{r['name']}</td><td align='right'>{fmt_int(r['z1'])}</td><td align='right'>{fmt_int(r['z2'])}</td></tr>"
    return html + "</tbody></table>"

def build_alert_html(alerts, vn_cur, hk_cur, vn):
    msgs = []
    if "trigger_vn" in alerts: msgs.append(f"🟩 <b>越南柴油跌破阈值</b>：当前价 {fmt_int(vn_cur)} VND/L，已跌破 25000 VND/L")
    if "recover_vn" in alerts: msgs.append(f"🟥 <b>越南柴油回升</b>：当前价 {fmt_int(vn_cur)} VND/L，已回升超过 25000 VND/L")
    if "trigger_hk" in alerts: msgs.append(f"🟩 <b>香港柴油跌破阈值</b>：当前价 {fmt_num(hk_cur)} HKD/L，已跌破 37 HKD/L")
    if "recover_hk" in alerts: msgs.append(f"🟥 <b>香港柴油回升</b>：当前价 {fmt_num(hk_cur)} HKD/L，已回升超过 37 HKD/L")
    
    color = "#27ae60" if "trigger" in str(alerts) else "#c9302c"
    bg = "#eef9f1" if "trigger" in str(alerts) else "#fdf0ef"
    
    html = f"""<div style="background:{bg};border:2px solid {color};border-radius:8px;padding:16px 20px;margin-bottom:20px;"><div style="font-size:16px;font-weight:700;color:{color};">⚠️ 油价预警</div>"""
    for m in msgs: html += f"""<div style="font-size:14px;color:#1f2d3d;margin-top:8px;line-height:1.7;">{m}</div>"""
    html += "</div>"
    
    if vn: html += section("越南油价明细", render_vietnam(vn))
    return html

def main():
    if not MAIL_USERNAME or not MAIL_PASSWORD: raise SystemExit("缺少 MAIL_USERNAME 或 MAIL_PASSWORD")
    today = bj_now().strftime("%Y-%m-%d")

    dom, hk, intl, vn = fetch_domestic(), fetch_hongkong(), fetch_international(), fetch_vietnam()
    prov_h, hk_h, vn_h = load_province_history(), load_hk_history(), load_vn_history()
    cal, _ = load_price_calendar(), load_intl_monthly()

    # 获取当前价格
    vn_cur = next((r["z1"] for r in vn["rows"] if r["prefix"] == "DO005S_M2"), None) if vn else None
    hk_cur = next((r["retail"] for r in hk["rows"] if "柴油" in r["fuel"]), None) if hk else None

    # 预警判断
    state = {"vn_below": False, "hk_below": False, "last_alert_date": None}
    if os.path.exists(ALERT_STATE_FILE):
        try:
            with open(ALERT_STATE_FILE, encoding="utf-8") as f: state.update(json.load(f))
        except: pass
        
    alerts = check_alerts(vn_cur, hk_cur, state, today) if state.get("last_alert_date") != today else []
    print(f"[INFO] 越南={vn_cur}, 香港={hk_cur}, 预警={alerts}")

    # 发日报
    summary_html = build_summary(dom, hk, vn, prov_h, hk_h, vn_h, cal)
    daily_body = f"""<html><body style="font-family:Arial,'Microsoft YaHei',sans-serif;color:#1f2d3d;max-width:960px;margin:auto;padding:16px;"><h2 style="color:#c9302c;text-align:center;">每日油价汇总 · {today}</h2><p style="color:#8492a6;font-size:12px;text-align:center;">自动推送，数据仅供参考</p>{summary_html}{section("一、国内各省市油价", render_domestic(dom))}{section("二、香港油价", render_hongkong(hk))}{section("三、国际原油", render_international(intl))}{section("四、越南油价", render_vietnam(vn))}<hr style="border:none;border-top:1px dashed #e4e7ed;margin:24px 0;"><div style="text-align:center;margin:20px 0;"><a href="{BASE_URL}" style="display:inline-block;padding:10px 24px;background:#c9302c;color:#fff;text-decoration:none;border-radius:6px;font-size:14px;font-weight:600;">🌐 查看完整网页版</a></div></body></html>"""
    try:
        msg = MIMEMultipart("alternative"); msg["Subject"] = Header(f"【每日油价】{today}", "utf-8"); msg["From"] = MAIL_USERNAME; msg["To"] = ", ".join(DAILY_RECEIVERS); msg.attach(MIMEText(daily_body, "html", "utf-8"))
        with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as server:
            server.login(MAIL_USERNAME, MAIL_PASSWORD); server.sendmail(MAIL_USERNAME, DAILY_RECEIVERS, msg.as_string())
        print("✓ 日报已发送")
    except Exception as e: print("[ERROR] 日报发送失败:", e)

    # 发预警
    if alerts:
        alert_body = f"""<html><body style="font-family:Arial,'Microsoft YaHei',sans-serif;color:#1f2d3d;max-width:960px;margin:auto;padding:16px;"><h2 style="color:#c9302c;text-align:center;">油价预警 · {today}</h2>{build_alert_html(alerts, vn_cur, hk_cur, vn)}<hr style="border:none;border-top:1px dashed #e4e7ed;margin:24px 0;"><div style="text-align:center;margin:20px 0;"><a href="{BASE_URL}" style="display:inline-block;padding:10px 24px;background:#c9302c;color:#fff;text-decoration:none;border-radius:6px;font-size:14px;font-weight:600;">🌐 查看完整网页版</a></div></body></html>"""
        try:
            msg = MIMEMultipart("alternative"); msg["Subject"] = Header(f"⚠️ 油价预警 · {today}", "utf-8"); msg["From"] = MAIL_USERNAME; msg["To"] = ", ".join(ALERT_RECEIVERS); msg.attach(MIMEText(alert_body, "html", "utf-8"))
            with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=30) as server:
                server.login(MAIL_USERNAME, MAIL_PASSWORD); server.sendmail(MAIL_USERNAME, ALERT_RECEIVERS, msg.as_string())
            print("✓ 预警已发送")
        except Exception as e: print("[ERROR] 预警发送失败:", e)
        
        state["last_alert_date"] = today

    # 更新状态
    if vn_cur is not None: state["vn_below"] = (vn_cur < VN_DIESEL_THRESHOLD)
    if hk_cur is not None: state["hk_below"] = (hk_cur < HK_DIESEL_THRESHOLD)
    with open(ALERT_STATE_FILE, "w", encoding="utf-8") as f: json.dump(state, f, ensure_ascii=False, indent=2)

if __name__ == "__main__":
    main()
