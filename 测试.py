import requests
import pandas as pd
from datetime import datetime, timedelta
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# ================= 配置 =================
CALENDAR_FILE = r"C:\Users\31007384\Desktop\新建文件夹\油价调整日历.xlsx"
HISTORY_FILE  = r"C:\Users\31007384\Desktop\新建文件夹\各省历史油价数据.xlsx"
API_URL = "https://v2.xxapi.cn/api/oilPrice"

# 全国 31 个省级行政区（不含港澳台）
PROVINCES = [
    "北京市","天津市","河北省","山西省","辽宁省","吉林省","黑龙江省",
    "上海市","江苏省","浙江省","安徽省","福建省","江西省","山东省",
    "湖北省","湖南省","河南省","海南省","重庆市","广东省","广西壮族自治区",
    "宁夏","甘肃省","新疆",
    "内蒙古","四川省","贵州省","云南省","陕西省","青海省","西藏",
]

# 接口省份名 → 本地 sheet 名
API_TO_LOCAL = {
    "北京市":"北京市","天津市":"天津市","河北省":"河北省","山西省":"山西省",
    "辽宁省":"辽宁省","吉林省":"吉林省","黑龙江省":"黑龙江省","上海市":"上海市",
    "江苏省":"江苏省","浙江省":"浙江省","安徽省":"安徽省","福建省":"福建省",
    "江西省":"江西省","山东省":"山东省","湖北省":"湖北省","湖南省":"湖南省",
    "河南省":"河南省","海南省":"海南省","重庆市":"重庆市","广东省":"广东省",
    "广西":"广西壮族自治区","宁夏":"宁夏","甘肃省":"甘肃省","新疆":"新疆",
    "内蒙古":"内蒙古","四川省":"四川省","贵州省":"贵州省","云南省":"云南省",
    "陕西省":"陕西省","青海省":"青海省","西藏":"西藏",
}

FUEL_MAP = {"n89":"89号汽油","n92":"92号汽油","n95":"95号汽油","n98":"98号汽油","n0":"0号柴油"}


# ============ 1. 拉取接口数据 ============
def fetch_today_prices():
    r = requests.get(API_URL, timeout=15)
    r.raise_for_status()
    js = r.json()
    if js.get("code") != 200:
        raise RuntimeError(f"接口返回异常：{js.get('msg')}")

    result = {}
    api_date = None
    for item in js["data"]:
        api_date = item.get("date")
        region = item.get("regionName")
        local = API_TO_LOCAL.get(region)
        if not local:
            continue
        price = {cn: item.get(en) for en, cn in FUEL_MAP.items()}
        result[local] = price
    return result, api_date


# ============ 2. 写入/更新历史表 ============
def update_history(province, date_str, price_dict):
    try:
        df = pd.read_excel(HISTORY_FILE, sheet_name=province)
    except Exception:
        df = pd.DataFrame(columns=["日期","89号汽油","92号汽油","95号汽油","98号汽油","0号柴油"])

    df["日期"] = pd.to_datetime(df["日期"], errors="coerce").dt.strftime("%Y-%m-%d")
    df = df.dropna(subset=["日期"])
    target = pd.to_datetime(date_str).strftime("%Y-%m-%d")

    if target in df["日期"].values:
        print(f"⏸ {province} {target} 已存在，跳过")
        return False

    new_row = {"日期": target}
    for k, v in price_dict.items():
        new_row[k] = v if v not in (None, "") else None

    df_new = pd.concat([pd.DataFrame([new_row]), df], ignore_index=True)
    df_new = df_new.sort_values("日期", ascending=False)

    with pd.ExcelWriter(HISTORY_FILE, engine="openpyxl", mode="a",
                        if_sheet_exists="replace") as writer:
        df_new.to_excel(writer, sheet_name=province, index=False)
    print(f"✅ {province} 新增 {target}：{price_dict}")
    return True


# ============ 3. 回填日历表 ============
def fill_calendar():
    cal_wb = load_workbook(CALENDAR_FILE)
    cal_ws = cal_wb.worksheets[0]
    hist_wb = load_workbook(HISTORY_FILE)

    province_dates = {}
    for p in PROVINCES:
        if p in hist_wb.sheetnames:
            ws = hist_wb[p]
            dates = set()
            for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
                if row[0]:
                    d = row[0]
                    if isinstance(d, str):
                        try:
                            d = pd.to_datetime(d).date()
                        except Exception:
                            continue
                    elif isinstance(d, datetime):
                        d = d.date()
                    dates.add(d)
            province_dates[p] = dates
        else:
            province_dates[p] = set()

    header = [c.value for c in cal_ws[1]]
    province_col = {n: i+1 for i, n in enumerate(header) if n in PROVINCES}

    green = PatternFill("solid", fgColor="C6EFCE")
    gray  = PatternFill("solid", fgColor="D9D9D9")

    for row in range(2, cal_ws.max_row + 1):
        cell_date = cal_ws.cell(row=row, column=1).value
        if not cell_date:
            continue
        if isinstance(cell_date, str):
            try:
                cell_date = pd.to_datetime(cell_date).date()
            except Exception:
                continue
        elif isinstance(cell_date, datetime):
            cell_date = cell_date.date()

        target = cell_date + timedelta(days=1)

        for p, col in province_col.items():
            cell = cal_ws.cell(row=row, column=col)
            if target in province_dates[p]:
                cell.value = "√"
                cell.fill = green
            else:
                cell.value = "未调整"
                cell.fill = gray

    cal_wb.save(CALENDAR_FILE)
    print("✅ 日历表回填完成")


# ============ 4. 查询任意日期油价 ============
def query_price(province, query_date, fuel_type="92号汽油"):
    if isinstance(query_date, str):
        query_date = pd.to_datetime(query_date).date()
    elif isinstance(query_date, datetime):
        query_date = query_date.date()

    df = pd.read_excel(HISTORY_FILE, sheet_name=province)
    df["日期"] = pd.to_datetime(df["日期"], errors="coerce").dt.date
    df = df.dropna(subset=["日期"]).sort_values("日期")

    valid = df[df["日期"] <= query_date]
    if valid.empty:
        return None
    return valid.iloc[-1][fuel_type]


# ============ 5. 每日主流程 ============
def daily_job():
    print("→ 拉取接口数据...")
    prices, api_date = fetch_today_prices()
    print(f"→ 接口日期：{api_date}，覆盖省份数：{len(prices)}")

    updated_any = False
    for province, price in prices.items():
        if province not in PROVINCES:
            continue
        if update_history(province, api_date, price):
            updated_any = True

    fill_calendar()
    print("→ 完成。是否新增过数据：", updated_any)


if __name__ == "__main__":
    daily_job()
    print("8/5 广东省 89号：", query_price("广东省", "2026-08-05", "89号汽油"))

    # 一键导出网页数据
    from export_json import export
    export()