# add_city_columns.py —— 往油价调整日历.xlsx 追加 17 个市级列
from openpyxl import load_workbook

CALENDAR_FILE = r"C:\Users\31007384\Desktop\新建文件夹\油价调整日历.xlsx"

CITY_COLS = [
    "广东省-深圳市","辽宁省-大连市","山东省-青岛市","福建省-厦门市",
    "浙江省-宁波市","西藏-拉萨市","新疆-乌鲁木齐市","新疆-克拉玛依市",
    "内蒙古-呼和浩特市","四川省-甘孜藏族自治州","四川省-阿坝藏族羌族自治州",
    "四川省-凉山彝族自治州","云南省-迪庆藏族自治州","云南省-怒江傈僳族自治州",
    "青海省-玉树藏族自治州","青海省-果洛藏族自治州","甘肃省-甘南藏族自治州",
]

wb = load_workbook(CALENDAR_FILE)
ws = wb.worksheets[0]

# 读表头，找最后一个非空列
headers = [c.value for c in ws[1]]
last_col = 1
for i, v in enumerate(headers, start=1):
    if v:
        last_col = i

print(f"当前表头最后列：{last_col}（{headers[last_col-1]}）")

# 检查是否已经加过
added = []
for name in CITY_COLS:
    if name in headers:
        print(f"  已存在，跳过：{name}")
        continue
    last_col += 1
    ws.cell(row=1, column=last_col).value = name
    added.append(name)

wb.save(CALENDAR_FILE)
print(f"✅ 新增 {len(added)} 列：")
for a in added:
    print("   ", a)